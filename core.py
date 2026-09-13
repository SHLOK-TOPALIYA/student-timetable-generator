import csv
import io
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pdfplumber
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
DAY_NORMALIZED = {d.upper(): d for d in DAYS}
DAY_NORMALIZED["THRUSDAY"] = "Thursday"

@dataclass(frozen=True)
class Entry:
    code: str
    section: str
    day: str
    time: str
    details: str = ""
    source: str = ""
    kind: str = "lecture"       # lecture / lab / tutorial
    group: str = ""             # lab/tutorial group, if applicable
    activity_code: str = ""      # code printed in the activity timetable
    association: str = "direct"  # direct / inferred / explicit

    @property
    def display_code(self):
        return f"{self.code} ({self.section})" if self.section else self.code

    @property
    def display_activity(self):
        label = "Lecture" if self.kind == "lecture" else self.kind.title()
        bits = [label]
        if self.group:
            bits.append(self.group)
        return " • ".join(bits)


def clean(s):
    """Normalize timetable text, including common PDF extraction artifacts."""
    text = str(s or "").replace("\xa0", " ")
    # pdfplumber can expose ReportLab bullet glyphs as literal (cid:127).
    text = re.sub(r"\(cid:\d+\)", " • ", text, flags=re.I)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def norm_code(s):
    s = clean(s).upper()
    s = re.sub(r"\s*\(\s*", "(", s)
    s = re.sub(r"\s*\)\s*", ")", s)
    return s


def split_code_section(raw):
    s = norm_code(raw)
    m = re.match(r"^(.*?)\s*\(\s*([A-Z])\s*\)$", s)
    if m:
        return clean(m.group(1)), m.group(2)
    m = re.match(r"^([A-Z0-9][A-Z0-9._-]*)\s*[- ]\s*([A-Z])$", s)
    if m:
        return m.group(1), m.group(2)
    return s, ""


SECTION_IN_TEXT_RE = re.compile(r"\(\s*SEC(?:TION)?\.?\s*-?\s*([A-Z0-9]{1,3})\s*\)", re.I)
CREDIT_PATTERN_RE = re.compile(r"^\d+-\d+-\d+-\d+$")


def extract_section_from_text(text):
    """Pull a 'Sec A' / 'Section B' style marker out of a course title.

    Some university timetables print the section inside the course-title
    cell (e.g. 'Data Structures (Sec A)') instead of next to the course
    code. Returns (cleaned_text, section) where section is '' if none found.
    """
    s = clean(text)
    if not s:
        return s, ""
    m = SECTION_IN_TEXT_RE.search(s)
    if not m:
        return s, ""
    section = m.group(1).upper()
    cleaned = clean(SECTION_IN_TEXT_RE.sub("", s))
    return cleaned, section


def looks_like_course(value):
    s = clean(value)
    if not s or s in {"-", ".", "—", "_"}:
        return False
    # Accept normal codes and named elective-style codes such as PC1(ICT&CS).
    return bool(re.match(r"^[A-Z]{1,10}\d{1,6}(?:\s*\([A-Z0-9& +_-]+\))?$", s.upper()))


def time_like(value):
    s = clean(value)
    return bool(re.match(r"^\d{1,2}:\d{2}\s*-\s*\d{1,2}:\d{2}", s)) or bool(
        re.match(r"^\d{1,2}:\d{2}\s*(AM|PM)$", s, re.I)
    )


def time_sort_key(t):
    s = clean(t).upper()
    m = re.search(r"(\d{1,2}):(\d{2})", s)
    if not m:
        return (99, 99, s)
    h = int(m.group(1))
    minute = int(m.group(2))
    if "PM" in s and h < 12:
        h += 12
    elif "AM" not in s and "PM" not in s and 1 <= h <= 6:
        h += 12
    return (h, minute, s)


def source_title_from_excel(ws):
    parts = []
    for r in range(1, min(ws.max_row, 12) + 1):
        v = clean(ws.cell(r, 1).value).replace("\n", " ")
        if v.upper() in {"TIME SLOT", "MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY"}:
            break
        if v and v not in parts:
            parts.append(v)
        if len(parts) >= 2:
            break
    return " — ".join(parts) if parts else "Student Timetable"


def find_day_columns(ws):
    best = None
    for r in range(1, min(ws.max_row, 40) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            v = clean(ws.cell(r, c).value).upper()
            if v in DAY_NORMALIZED:
                found[DAY_NORMALIZED[v]] = c
        if len(found) >= 5:
            best = (r, found)
            break
    if not best:
        raise ValueError("Could not find a Monday–Friday header in the Excel file.")
    header_row, starts = best
    return header_row, starts


def parse_lecture_excel(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    target = None
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(50, ws.max_row)):
            if sum(clean(c.value).upper() in DAY_NORMALIZED for c in row) >= 5:
                target = ws
                break
        if target:
            break
    if target is None:
        raise ValueError("No timetable sheet containing Monday–Friday was found.")
    ws = target
    header_row, starts = find_day_columns(ws)
    course_cols = {}
    for day, start in starts.items():
        course_col = start
        for offset in range(0, 5):
            v = clean(ws.cell(header_row + 1, start + offset).value).upper()
            if v == "COURSE":
                course_col = start + offset
                break
        course_cols[day] = course_col

    entries, current_time = [], None
    for r in range(header_row + 1, ws.max_row + 1):
        first = clean(ws.cell(r, 1).value)
        if time_like(first):
            current_time = first
            continue
        if not current_time:
            continue
        for day in DAYS:
            cc = course_cols[day]
            raw = clean(ws.cell(r, cc).value)
            if not looks_like_course(raw):
                continue
            name = clean(ws.cell(r, cc + 1).value)
            # Master/institute-wide timetables (Code | Title | Credit | Type |
            # Faculty | Room) don't have a "COURSE" sub-header, so the room
            # isn't always right next to the title. Detect that layout by
            # checking whether the next column looks like a credit pattern
            # (e.g. "3-0-0-3"); if so, the room is 3 columns further along.
            credit_like = CREDIT_PATTERN_RE.match(clean(ws.cell(r, cc + 2).value))
            room = clean(ws.cell(r, cc + 5).value) if credit_like else clean(ws.cell(r, cc + 2).value)
            code, section = split_code_section(raw)
            if not section:
                # Some formats embed the section in the title instead, e.g.
                # "Data Structures (Sec A)". Detect and strip it out.
                name, section = extract_section_from_text(name)
            details = " • ".join(x for x in [name, room] if x and x not in {"-", "."})
            entries.append(Entry(code, section, day, current_time, details, "Excel", "lecture", "", code))
    if not entries:
        raise ValueError("No lecture entries could be extracted from the Excel timetable.")
    # Master/institute-wide timetables often list the same shared lecture
    # under several student batches (e.g. the same course+section taught to
    # two programmes at once). That produces exact duplicate rows which must
    # be collapsed, otherwise they look like a course clashing with itself.
    entries = list(dict.fromkeys(entries))
    return entries, source_title_from_excel(ws)


def extract_group_tokens(cell):
    s = clean(cell).upper()
    if not s:
        return []
    # Handles G1/G2, A/B, L1/L2, Group A/B and Tut. G1.
    s = re.sub(r"\bGROUP\s*", "", s)
    s = re.sub(r"\bTUTORIAL\s*", "", s)
    tokens = re.findall(r"\b[A-Z]{1,3}\s*\d{0,2}\b", s)
    # Avoid treating ordinary words such as TUT as groups.
    out = []
    for tok in tokens:
        tok = clean(tok).replace(" ", "")
        if tok in {"TUT", "LAB"}:
            continue
        if re.match(r"^[A-Z]{1,3}\d{0,2}$", tok):
            out.append(tok)
    return list(dict.fromkeys(out))


def is_tutorial_cell(cell):
    return bool(re.search(r"\bTUT\.?\s*(?:GROUP\s*)?[A-Z]{1,3}\s*\d{0,2}\b", clean(cell), re.I))


def split_rooms(details):
    s = clean(details)
    if not s:
        return []
    # Prefer explicit separators used by institutional sheets.
    parts = re.split(r"\s*&\s*|\s*,\s*|\s+AND\s+", s, flags=re.I)
    return [clean(x).upper() for x in parts if clean(x)]


def map_groups_to_rooms(groups, rooms):
    if not groups:
        return {"": clean(" & ".join(rooms)) if rooms else ""}
    if len(rooms) == len(groups):
        return {g: rooms[i] for i, g in enumerate(groups)}
    # Common case: two group labels sharing a single lab room string.
    if len(rooms) == 1:
        return {g: rooms[0] for g in groups}
    return {g: " & ".join(rooms) for g in groups}


def parse_activity_excel(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    target = None
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(20, ws.max_row)):
            vals = {clean(c.value).upper() for c in row}
            if "LAB DETAILS" in vals and "COURSE" in vals and "MONDAY" in vals:
                target = ws
                break
        if target:
            break
    if target is None:
        raise ValueError("Could not find a Lab/Tutorial timetable sheet with Time Slot, Monday–Friday, Lab Details and Course columns.")

    ws = target
    header_row = None
    cols = {}
    for r in range(1, min(ws.max_row, 30) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            v = clean(ws.cell(r, c).value).upper()
            if v in {"TIME SLOT", "MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "LAB DETAILS", "COURSE", "FACULTY"}:
                found[v] = c
        if "MONDAY" in found and "FRIDAY" in found and "COURSE" in found:
            header_row = r
            cols = found
            break
    if header_row is None:
        raise ValueError("Could not identify the Lab/Tutorial timetable header.")

    day_cols = {DAY_NORMALIZED[d]: cols[d] for d in DAY_NORMALIZED if d in cols}
    time_col = cols.get("TIME SLOT", 1)
    room_col = cols.get("LAB DETAILS")
    course_col = cols.get("COURSE")
    entries = []
    current_time = None
    current_course = ""
    current_room_text = ""
    current_faculty = ""
    for r in range(header_row + 1, ws.max_row + 1):
        first = clean(ws.cell(r, time_col).value)
        if time_like(first):
            if clean(first) != clean(current_time):
                current_course = ""
                current_room_text = ""
                current_faculty = ""
            current_time = first
        if not current_time:
            continue
        raw_course = clean(ws.cell(r, course_col).value)
        room_text = clean(ws.cell(r, room_col).value) if room_col else ""
        faculty = clean(ws.cell(r, cols.get("FACULTY", course_col + 1)).value) if cols.get("FACULTY") else ""
        if looks_like_course(raw_course):
            current_course = raw_course
            current_room_text = room_text
            current_faculty = faculty
        elif raw_course:
            # Non-course labels do not replace the current activity.
            pass
        if not current_course:
            continue
        raw_course = current_course
        room_text = room_text or current_room_text
        faculty = faculty or current_faculty
        rooms = split_rooms(room_text)
        for day in DAYS:
            if day not in day_cols:
                continue
            cell = clean(ws.cell(r, day_cols[day]).value)
            if not cell or cell in {"-", "—", "_"}:
                continue
            tutorial = is_tutorial_cell(cell)
            groups = extract_group_tokens(cell)
            # Only treat group tokens as actual group labels when the cell is a group schedule.
            # Plain course-code cells are ungrouped activities.
            if not tutorial and raw_course.upper() in cell.upper():
                groups = []
            kind = "tutorial" if tutorial else "lab"
            mapping = map_groups_to_rooms(groups, rooms)
            code, section = split_code_section(raw_course)
            for group in (groups or [""]):
                room = mapping.get(group, room_text)
                details = " • ".join(x for x in [faculty, room] if x)
                entries.append(Entry(code, section, day, current_time, details, "Excel", kind, group, code))
    if not entries:
        raise ValueError("No Lab/Tutorial entries could be extracted from the uploaded workbook.")
    entries = list(dict.fromkeys(entries))
    return entries, source_title_from_excel(ws)


def extract_pdf_tables(data):
    tables = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            page_tables = page.extract_tables(table_settings={
                "vertical_strategy": "lines", "horizontal_strategy": "lines",
                "snap_tolerance": 4, "join_tolerance": 4, "intersection_tolerance": 5,
            }) or page.extract_tables()
            for t in page_tables or []:
                tables.append((page_no, t))
    return tables


def parse_lecture_pdf(data):
    entries = []
    title = "Student Timetable"
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if pdf.pages:
            text = pdf.pages[0].extract_text() or ""
            first = next((clean(x) for x in text.splitlines() if clean(x)), "")
            if first and "MONDAY" not in first.upper() and "TUESDAY" not in first.upper():
                title = first
    for page_no, table in extract_pdf_tables(data):
        if not table or not table[0]:
            continue
        header_idx, day_cols = None, {}
        for i, row in enumerate(table[:12]):
            found = {}
            for j, cell in enumerate(row):
                v = clean(cell).upper()
                if v in DAY_NORMALIZED:
                    found[DAY_NORMALIZED[v]] = j
            if len(found) >= 5:
                header_idx, day_cols = i, found
                break
        if header_idx is None:
            continue
        # Our generated PDF and many compact timetables: TIME | Mon..Fri.
        for r in range(header_idx + 1, len(table)):
            row = table[r] or []
            if not row:
                continue
            time = next((clean(c) for c in row[:2] if time_like(c)), None)
            if not time:
                continue
            for day in DAYS:
                j = day_cols.get(day)
                if j is None or j >= len(row):
                    continue
                cell = clean(row[j])
                if not cell or cell in {"—", "-", "_"}:
                    continue
                # Option 4 is for an existing lecture timetable. If a user
                # accidentally uploads a complete timetable PDF, do not turn
                # its LAB/TUTORIAL cells into duplicate lecture entries.
                cell_upper = cell.upper()
                if ("LAB" in cell_upper or "TUTORIAL" in cell_upper) and "LECTURE" not in cell_upper:
                    continue
                # Multiple entries in a cell are separated by line breaks.
                lines = [clean(x) for x in re.split(r"\n+", cell) if clean(x)]
                code_matches = []
                for line in lines:
                    m = re.match(r"^([A-Z]{1,10}\d{1,6}(?:\s*\([A-Z]\))?)(?=\s|$)", line, re.I)
                    if m:
                        code_matches.append((m.group(1), clean(line[len(m.group(1)):]).strip()))
                if not code_matches:
                    continue
                for raw_code, rest in code_matches:
                    code, section = split_code_section(raw_code)
                    # When reading a PDF generated by this app, the cell already
                    # contains the word LECTURE. The Entry kind supplies that label
                    # during PDF generation, so remove it from the stored details to
                    # avoid displaying "LECTURE" twice after an option-4 round trip.
                    rest = re.sub(r"^LECTURE\b\s*", "", rest, flags=re.I).strip()
                    if not section:
                        rest, section = extract_section_from_text(rest)
                    entries.append(Entry(code, section, day, time, rest, f"PDF p.{page_no}", "lecture", "", code))
    if not entries:
        raise ValueError("Could not read a lecture timetable from this PDF. Please use a text/table PDF or the original Excel file.")
    entries = list(dict.fromkeys(entries))
    return entries, title


def parse_csv(data):
    text = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise ValueError("CSV file is empty.")
    # Simple timetable CSV support: Time, Monday, Tuesday, Wednesday, Thursday, Friday.
    header = [clean(x).upper() for x in rows[0]]
    if not all(d in header for d in ["MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY"]):
        raise ValueError("CSV must contain Monday–Friday columns in the first row.")
    entries, current_time = [], None
    for row in rows[1:]:
        if not row:
            continue
        first = clean(row[0]) if row else ""
        if time_like(first):
            current_time = first
        if not current_time:
            continue
        for day in DAYS:
            j = header.index(day)
            if j >= len(row):
                continue
            raw = clean(row[j])
            if not raw or raw in {"-", "—"}:
                continue
            m = re.match(r"^([A-Z]{1,10}\d{1,6}(?:\s*\([A-Z]\))?)(?=\s|$)(.*)$", raw, re.I)
            if not m:
                continue
            code, section = split_code_section(m.group(1))
            rest = clean(m.group(2))
            if not section:
                rest, section = extract_section_from_text(rest)
            entries.append(Entry(code, section, day, current_time, rest, "CSV", "lecture", "", code))
    if not entries:
        raise ValueError("No timetable entries could be extracted from the CSV.")
    return entries, "Uploaded Timetable"


def parse_lecture_file(filename, data):
    ext = Path(filename).suffix.lower()
    if ext in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return parse_lecture_excel(data)
    if ext == ".pdf":
        return parse_lecture_pdf(data)
    if ext == ".csv":
        return parse_csv(data)
    raise ValueError("Supported lecture formats: XLSX, XLSM, PDF and CSV.")


def parse_activity_file(filename, data):
    ext = Path(filename).suffix.lower()
    if ext in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return parse_activity_excel(data)
    if ext == ".pdf":
        # A PDF Lab/Tutorial timetable can use the same compact table extraction.
        return parse_activity_pdf(data)
    if ext == ".csv":
        return parse_activity_csv(data)
    raise ValueError("Supported Lab/Tutorial formats: XLSX, XLSM, PDF and CSV.")


def parse_activity_pdf(data):
    # Reuse table extraction, looking for Lab Details / Course columns or compact group cells.
    entries = []
    title = "Lab/Tutorial Timetable"
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if pdf.pages:
            text = pdf.pages[0].extract_text() or ""
            first = next((clean(x) for x in text.splitlines() if clean(x)), "")
            if first:
                title = first
    for page_no, table in extract_pdf_tables(data):
        if not table:
            continue
        header = None
        for i, row in enumerate(table[:15]):
            vals = [clean(c).upper() for c in row]
            if "MONDAY" in vals and "COURSE" in vals:
                header = i
                break
        if header is None:
            continue
        row0 = table[header]
        day_cols = {DAY_NORMALIZED[clean(c).upper()]: j for j, c in enumerate(row0) if clean(c).upper() in DAY_NORMALIZED}
        course_col = next((j for j,c in enumerate(row0) if clean(c).upper()=="COURSE"), None)
        room_col = next((j for j,c in enumerate(row0) if clean(c).upper() in {"LAB DETAILS","ROOM","VENUE"}), None)
        if course_col is None:
            continue
        current_time = None
        for r in range(header+1,len(table)):
            row=table[r] or []
            if not row: continue
            current_time=next((clean(c) for c in row[:3] if time_like(c)), current_time)
            if not current_time or course_col>=len(row): continue
            raw=clean(row[course_col])
            if not looks_like_course(raw): continue
            room_text=clean(row[room_col]) if room_col is not None and room_col<len(row) else ""
            rooms=split_rooms(room_text)
            code,section=split_code_section(raw)
            for day in DAYS:
                j=day_cols.get(day)
                if j is None or j>=len(row): continue
                cell=clean(row[j])
                if not cell or cell in {"-","—"}: continue
                tut=is_tutorial_cell(cell)
                groups=extract_group_tokens(cell)
                if not tut and raw.upper() in cell.upper(): groups=[]
                kind="tutorial" if tut else "lab"
                mapping=map_groups_to_rooms(groups,rooms)
                for group in groups or [""]:
                    room=mapping.get(group,room_text)
                    entries.append(Entry(code,section,day,current_time,room,f"PDF p.{page_no}",kind,group,code))
    if not entries:
        raise ValueError("Could not read a Lab/Tutorial timetable from this PDF. Please use a text/table PDF or Excel.")
    entries = list(dict.fromkeys(entries))
    return entries,title


def parse_activity_csv(data):
    # Treat CSV like a standard activity table with Time, Mon..Fri, Course, Room columns.
    rows=list(csv.reader(io.StringIO(data.decode('utf-8-sig',errors='replace'))))
    if not rows: raise ValueError('CSV file is empty.')
    header=[clean(x).upper() for x in rows[0]]
    if 'COURSE' not in header or 'MONDAY' not in header: raise ValueError('Activity CSV must contain Course and Monday–Friday columns.')
    cc=header.index('COURSE'); rc=next((header.index(x) for x in ['LAB DETAILS','ROOM','VENUE'] if x in header),None)
    entries=[]; current_time=None
    for row in rows[1:]:
        if row and time_like(row[0]): current_time=clean(row[0])
        if not current_time or cc>=len(row): continue
        raw=clean(row[cc])
        if not looks_like_course(raw): continue
        code,section=split_code_section(raw); room=clean(row[rc]) if rc is not None and rc<len(row) else ''
        for day in DAYS:
            j=header.index(day)
            if j>=len(row): continue
            cell=clean(row[j])
            if not cell: continue
            tut=is_tutorial_cell(cell); groups=extract_group_tokens(cell)
            if not tut and raw.upper() in cell.upper(): groups=[]
            kind='tutorial' if tut else 'lab'
            for group in groups or ['']:
                entries.append(Entry(code,section,day,current_time,room,'CSV',kind,group,code))
    if not entries: raise ValueError('No Lab/Tutorial entries could be extracted from the CSV.')
    return list(dict.fromkeys(entries)),'Uploaded Lab/Tutorial Timetable'


def course_catalog(entries):
    out=defaultdict(list)
    for e in entries: out[e.code].append(e)
    return out


def section_options(entries, code):
    return sorted({e.section for e in entries if e.code==code and e.section})


def extract_numeric_code(code):
    m=re.match(r'^([A-Z]+)(\d+)$', norm_code(code))
    return (m.group(1),int(m.group(2))) if m else None


def find_activity_associations(lecture_code, activity_entries, explicit_map=None):
    """Return activity entries only for an exact code or an explicit user-approved map.

    IMPORTANT: Similar-looking course codes are never treated as proof that two
    courses are the same. A separate activity code such as CT304 may correspond
    to lecture CT303 at one university, but CT102 may be a completely different
    course from CT101 at another. The timetable source and/or explicit user
    confirmation must establish the relationship.
    """
    code=norm_code(lecture_code)
    mapped=norm_code((explicit_map or {}).get(code, ''))
    target=mapped or code
    direct=[e for e in activity_entries if e.code==target]
    if direct:
        return direct, 'explicit' if mapped else 'direct'
    return [], 'none'


def candidate_activity_codes(lecture_code, activity_entries):
    """Suggest possible activity codes for user confirmation only.

    These are suggestions, not automatic matches. Same-prefix/near-number codes
    are surfaced because they are useful for cases such as CT303 -> CT304, but
    the user must explicitly approve the mapping.
    """
    code=norm_code(lecture_code)
    parsed=extract_numeric_code(code)
    codes=sorted({e.code for e in activity_entries})
    if not parsed:
        return codes
    prefix,num=parsed
    same_prefix=[]
    for c in codes:
        p=extract_numeric_code(c)
        if p and p[0]==prefix:
            same_prefix.append((abs(p[1]-num), p[1], c))
    same_prefix.sort()
    return [c for _,_,c in same_prefix[:8]]


def select_lecture_entries(entries, requested):
    result=[]
    for code,section in requested:
        for e in entries:
            if e.code!=code: continue
            if section:
                if e.section==section: result.append(e)
            elif not e.section:
                result.append(e)
    return result


def select_activity_entries(entries, requested_codes, group_by_kind=None, explicit_map=None):
    group_by_kind=group_by_kind or {}
    result=[]
    for target_code in requested_codes:
        assoc,mode=find_activity_associations(target_code,entries,explicit_map)
        if not assoc: continue
        selected_groups=group_by_kind.get(target_code,{})
        for e in assoc:
            if e.group:
                g=selected_groups.get(e.kind)
                if g and e.group!=g: continue
            result.append(Entry(e.code,e.section,e.day,e.time,e.details,e.source,e.kind,e.group,target_code,mode))
    return result


def unique_activity_codes(entries, selected_codes, explicit_map=None):
    result=[]
    for code in selected_codes:
        assoc,_=find_activity_associations(code,entries,explicit_map)
        if assoc: result.append(code)
    return result


def group_options_for_codes(activity_entries, selected_codes, kind, explicit_map=None):
    sets=[]
    for code in selected_codes:
        assoc,_=find_activity_associations(code,activity_entries,explicit_map)
        opts={e.group for e in assoc if e.kind==kind and e.group}
        if opts: sets.append(opts)
    if not sets: return []
    common=set.intersection(*sets)
    return sorted(common) if common else sorted(set.union(*sets))


def conflict_groups(entries):
    by=defaultdict(list)
    for e in entries: by[(e.day,e.time)].append(e)
    return [(k,v) for k,v in by.items() if len(v)>1]


def parse_duration_minutes(t):
    m=re.match(r'^(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})',clean(t))
    if not m: return None
    h1,mi1,h2,mi2=map(int,m.groups())
    if h2<h1: h2+=12
    return (h1*60+mi1,h2*60+mi2)


def overlap(a,b):
    x=parse_duration_minutes(a); y=parse_duration_minutes(b)
    return bool(x and y and max(x[0],y[0])<min(x[1],y[1]))


def detect_conflicts(selected):
    by_day=defaultdict(list)
    for e in selected: by_day[e.day].append(e)
    conflicts=[]
    for day,items in by_day.items():
        for i in range(len(items)):
            for j in range(i+1,len(items)):
                if overlap(items[i].time,items[j].time):
                    conflicts.append(((day,f"{items[i].time} ↔ {items[j].time}"),[items[i],items[j]]))
    return conflicts


def safe_filename(title):
    s=clean(title)
    s=re.sub(r'[^A-Za-z0-9 _-]+','',s)
    s=re.sub(r'\s+','_',s).strip('_')
    return (s[:120] or 'My_Timetable')+'.pdf'


def escape_html(s):
    return str(s or '').replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')


def make_pdf(selected, title='Student Timetable'):
    buffer=io.BytesIO()
    doc=SimpleDocTemplate(buffer,pagesize=landscape(A4),rightMargin=8*mm,leftMargin=8*mm,topMargin=7*mm,bottomMargin=7*mm,title=title,author='Student Timetable Generator')
    times=sorted({e.time for e in selected},key=time_sort_key)
    styles=getSampleStyleSheet()
    title_style=ParagraphStyle('T',parent=styles['Title'],fontName='Helvetica-Bold',fontSize=15,leading=18,alignment=TA_CENTER,spaceAfter=5)
    head=ParagraphStyle('H',parent=styles['Normal'],fontName='Helvetica-Bold',fontSize=8.5,leading=10,alignment=TA_CENTER)
    cell=ParagraphStyle('C',parent=styles['Normal'],fontName='Helvetica',fontSize=7,leading=8.5,alignment=TA_CENTER)
    time_style=ParagraphStyle('TM',parent=cell,fontName='Helvetica-Bold')
    legend=ParagraphStyle('L',parent=styles['Normal'],fontSize=7.5,leading=9,alignment=TA_LEFT)
    story=[Paragraph(escape_html(title),title_style)]
    data=[[Paragraph('TIME',head)]+[Paragraph(d.upper(),head) for d in DAYS]]
    for t in times:
        row=[Paragraph(escape_html(t),time_style)]
        for day in DAYS:
            items=sorted([e for e in selected if e.day==day and e.time==t],key=lambda e:(e.kind,e.code,e.group,e.section))
            if not items:
                row.append(Paragraph('—',cell)); continue
            blocks=[]
            for e in items:
                label=e.code
                if e.section: label+=f' ({e.section})'
                tag=e.kind.upper()
                if e.group: tag+=f' • {e.group}'
                txt=f'<b>{escape_html(label)}</b><br/><font size="6.5">{escape_html(tag)}</font>'
                if e.details: txt+=f'<br/><font size="6.5">{escape_html(e.details)}</font>'
                blocks.append(Paragraph(txt,cell))
            row.append(Spacer(1,0) if False else Paragraph('<br/><br/>'.join([b.getPlainText() if False else '' for b in []]),cell) if False else Paragraph('<br/><br/>'.join([f'<b>{escape_html(e.code + (" ("+e.section+")" if e.section else ""))}</b><br/><font size="6.5">{escape_html((e.kind.upper() + (" • "+e.group if e.group else "")))}</font>' + (f'<br/><font size="6.5">{escape_html(e.details)}</font>' if e.details else '') for e in items]),cell))
        data.append(row)
    table=Table(data,colWidths=[29*mm]+[50*mm]*5,repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#F7D774')),
        ('BACKGROUND',(0,1),(0,-1),colors.HexColor('#F5E6A8')),
        ('BACKGROUND',(1,1),(-1,-1),colors.HexColor('#EAF1F8')),
        ('GRID',(0,0),(-1,-1),0.65,colors.HexColor('#666666')),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(0,0),(-1,-1),'CENTER'),
        ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),
        ('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3),
    ]))
    story.append(table)
    story.append(Spacer(1,5*mm))
    rows=[[Paragraph('<b>COURSE</b>',head),Paragraph('<b>TYPE</b>',head),Paragraph('<b>GROUP</b>',head),Paragraph('<b>SECTION</b>',head),Paragraph('<b>DETAILS</b>',head)]]
    seen=set()
    for e in sorted(selected,key=lambda x:(x.code,x.kind,x.group,x.section,x.day,x.time)):
        key=(e.code,e.kind,e.group,e.section,e.details)
        if key in seen: continue
        seen.add(key)
        rows.append([Paragraph(escape_html(e.code),legend),Paragraph(escape_html(e.kind.title()),legend),Paragraph(escape_html(e.group or '—'),legend),Paragraph(escape_html(e.section or '—'),legend),Paragraph(escape_html(e.details or '—'),legend)])
    if len(rows)>1:
        lg=Table(rows,colWidths=[32*mm,28*mm,25*mm,25*mm,100*mm])
        lg.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#F6E2D5')),('GRID',(0,0),(-1,-1),0.55,colors.HexColor('#666666')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),2.5),('BOTTOMPADDING',(0,0),(-1,-1),2.5),('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3)]))
        story.append(lg)
    doc.build(story)
    return buffer.getvalue()
